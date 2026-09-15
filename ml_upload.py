"""
Sube publicaciones a MercadoLibre a partir del Excel generado por scraper.py.

Uso:
    python ml_upload.py --dry-run --limit 5      # probar sin publicar nada
    python ml_upload.py --in productos_rodo.xlsx --out resultado_publicacion.xlsx

Ver MANUAL.md para la explicacion completa de como funciona cada pieza,
que significa cada columna del excel de resultado, y como operar todo esto
a mano si hiciera falta.
"""
import argparse
import glob
import io
import json
import os
import re
import sys
import time
import unicodedata

import requests
from openpyxl import load_workbook, Workbook
from PIL import Image, ImageFilter

from ml_auth import get_valid_token

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

API_BASE = "https://api.mercadolibre.com"
# Publicamos con mas de 1 unidad de stock a proposito: si vendemos 1 y
# quedamos en 0, ML pausa la publicacion sola (no deja vender sin stock) y
# hay que reponerla a mano. Con margen de stock, una venta no pausa nada.
DEFAULT_STOCK = 3
API_TIMEOUT = 30
MAX_ATTRIBUTE_RETRIES = 10


class TimeoutSession(requests.Session):
    """Sesion con timeout por default en todos los pedidos -- sin esto, un
    corte de red puede dejar la conexion colgada para siempre (paso en
    produccion: 'timeout=None' + un corte de red tumbo todo el script a
    mitad de una tanda grande)."""

    def request(self, *args, **kwargs):
        kwargs.setdefault("timeout", API_TIMEOUT)
        return super().request(*args, **kwargs)

# (alto_cm, ancho_cm, largo_cm, peso_g) -- estimaciones conservadoras, no exactas.
# Revisar a mano los productos grandes (heladeras, lavarropas, aires) antes de
# confiar en el costo de envio que calcule ML.
PACKAGE_DEFAULT = (20, 20, 20, 1000)

PACKAGE_KEYWORDS = [
    (("heladera", "freezer", "refrigerador"), (160, 70, 70, 60000)),
    (("lavarropas", "lavaseca", "secarropas"), (90, 60, 65, 40000)),
    (("aire acondicionado", "split", "acondicionado"), (35, 80, 30, 15000)),
    (("lavavajillas",), (85, 60, 60, 35000)),
    (("cocina", "horno empotrable", "anafe"), (90, 60, 60, 30000)),
    (("tv", "smart tv", "televisor", "qled"), (10, 100, 60, 9000)),
    (("notebook", "laptop"), (5, 35, 25, 2200)),
    (("celular", "smartphone", "iphone"), (5, 16, 8, 300)),
    (("tablet",), (8, 25, 18, 600)),
    (("microondas",), (30, 45, 35, 12000)),
    (("licuadora", "batidora", "mixer"), (25, 20, 20, 2000)),
    (("cafetera",), (30, 25, 20, 2500)),
    (("impresora",), (30, 40, 40, 5000)),
    (("consola",), (18, 40, 30, 3000)),
    (("auricular", "parlante", "speaker"), (15, 20, 15, 800)),
]

# Palabras clave para decidir el markup "grande" (electrodomesticos de
# linea blanca grandes) vs "chico" (todo lo demas) -- ver seccion 8 de
# MANUAL.md, tabla de precios por cuotas del 3/11/25. El precio publicado
# es siempre el de "1 pago"; los recargos por cuotas de esa tabla se
# manejan aparte en la configuracion de cuotas/financiacion de ML, no en
# el precio del articulo (ML solo permite un precio unico por publicacion).
LARGE_ITEM_KEYWORDS = ("heladera", "freezer", "refrigerador", "lavarropas", "lavaseca",
                       "secarropas", "aire acondicionado", "split", "acondicionado",
                       "lavavajillas", "cocina")


def pick_markup(title, markup_chico, markup_grande):
    text = strip_accents((title or "").lower())
    for kw in LARGE_ITEM_KEYWORDS:
        if re.search(rf"\b{re.escape(kw)}\b", text):
            return markup_grande
    return markup_chico


CATEGORY_ESTIMATES = {
    "celulares": (5, 16, 8, 300),
    "telefonia": (8, 20, 15, 500),
    "imagen y sonido": (15, 60, 40, 8000),
    "electrodomesticos": (35, 40, 35, 6000),
    "electro hogar": (25, 30, 25, 3000),
    "informatica": (10, 40, 30, 2500),
    "climatizacion": (40, 70, 30, 12000),
    "cuidado personal": (10, 15, 10, 500),
    "deporte y aire libre": (25, 30, 20, 2000),
}

# Valores genericos de ultimo recurso para atributos que ML pide y Rodo no
# publica. Son inferencias razonables (ej: en Argentina casi todo es 220V),
# no datos reales del producto puntual -- conviene revisarlos si el producto
# es una excepcion conocida.
GENERIC_ATTRIBUTE_DEFAULTS = {
    "voltaje": "220V",
    "genero": "Unisex",
    "tipo alimentacion": "Pilas",
    "tipos alimentacion": "Pilas",
    "coleccion": "Sin coleccion",
    "temporada": "Todo el año",
    "pais fabricacion": "China",
    "requiere ensamblado": "Sí",
    "es plegable": "No",
    "plegable": "No",
    "incluye manual ensamblado": "Sí",
    "incluye herramientas ensamblado": "No",
    "es minibar": "No",
}

# Sinonimos: ML a veces pide un atributo con un nombre distinto al que usa
# Rodo para (casi) lo mismo. Cada entrada dice, para el nombre normalizado
# que pide ML, en que otras etiquetas normalizadas de Rodo buscar el valor.
ATTRIBUTE_ALIASES = {
    "linea procesador": ["tipo procesador", "version procesador", "procesador"],
    "memoria ram": ["cantidad memoria", "memoria"],
    "capacidad almacenamiento": ["capacidad disco rigido", "almacenamiento"],
    "tamano pantalla": ["cantidad pulgadas", "tam pantalla", "pulgadas"],
}

# Palabras clave en el titulo -> valor de "Plataforma" para videojuegos.
PLATFORM_KEYWORDS = [
    ("ps5", "PlayStation 5"), ("playstation 5", "PlayStation 5"),
    ("ps4", "PlayStation 4"), ("playstation 4", "PlayStation 4"),
    ("xbox series", "Xbox Series X/S"), ("xbox one", "Xbox One"),
    ("nintendo switch 2", "Nintendo Switch 2"), ("nintendo switch", "Nintendo Switch"),
    ("pc", "PC"),
]

PUBLISHED_STATE_FILE = "published_skus.json"

# Palabras que, si aparecen en el nombre del producto de catalogo pero NO en
# nuestro titulo, sugieren que es un combo/bundle distinto al que vende Rodo
# (ej: "con resma de papel", "kit", "+ funda") -> no conviene pegarse a esa ficha.
BUNDLE_RISK_WORDS = {"resma", "combo", "kit", "pack", "bundle", "regalo", "gratis",
                     "oferta", "promo", "accesorio", "funda", "case", "cargador"}

BRAND_ATTR_LINE = re.compile(r"^\s*Marca\s*:\s*(.+)$", re.IGNORECASE | re.MULTILINE)
MISSING_ATTR_MSG = re.compile(r'El campo "(.+?)" es obligatorio')
MISSING_ATTR_IDS_MSG = re.compile(r"attributes \[([A-Z0-9_,\s]+)\] are (?:all )?required")
INVALID_ATTR_ID_MSG = re.compile(r"Attribute \[([A-Z0-9_]+)\] is not valid")
GTIN_INVALID_MSG = re.compile(r"Product Identifier \[GTIN\]")

_category_attrs_cache = {}


PACKAGING_BUFFER_CM = 5  # la caja de envio es un poco mas grande que el producto solo

MEASURE_RE_TEMPLATE = r"{label}\s*:?\s*([\d]+[.,]?\d*)(?:\s*-\s*([\d]+[.,]?\d*))?\s*cm"


def _extract_max_measurement(text, label):
    """Busca todas las menciones de ese lado (Alto/Ancho/Profundidad) en el
    texto y devuelve la mas grande -- cubre casos con rangos ('40 - 80 cm')
    o mas de una medida (ej. aires con unidad interior Y exterior), sin
    correr el riesgo de quedarse corto."""
    pattern = re.compile(MEASURE_RE_TEMPLATE.format(label=label), re.IGNORECASE)
    values = []
    for m in pattern.finditer(text or ""):
        for g in m.groups():
            if g:
                values.append(float(g.replace(",", ".")))
    return max(values) if values else None


def parse_rodo_dimensions(description):
    """Lee las medidas REALES del producto desde la ficha tecnica de Rodo
    (linea 'Medidas: Alto: X cm x Ancho: Y cm x Profundidad: Z cm'), en vez
    de adivinar por categoria/palabra clave. Son medidas del PRODUCTO, no
    del paquete de envio -- por eso se les suma PACKAGING_BUFFER_CM.
    Se detecto una infraccion real de ML por declarar un paquete mas chico
    que el producto real, asi que esto tiene prioridad sobre la estimacion
    generica siempre que este disponible."""
    if not description:
        return None
    h = _extract_max_measurement(description, "Alto")
    w = _extract_max_measurement(description, "Ancho")
    d = _extract_max_measurement(description, "Profundidad")
    if h and w and d:
        return (round(h + PACKAGING_BUFFER_CM), round(w + PACKAGING_BUFFER_CM), round(d + PACKAGING_BUFFER_CM))
    return None


DIMENSION_SANITY_RANGE_CM = (2, 250)
WEIGHT_SANITY_RANGE_G = (20, 200000)


def _parse_cm(value_name):
    if not value_name:
        return None
    m = re.search(r"([\d.,]+)\s*(cm|mm|m)\b", value_name, re.IGNORECASE)
    if not m:
        return None
    num = float(m.group(1).replace(",", "."))
    unit = m.group(2).lower()
    if unit == "m":
        num *= 100
    elif unit == "mm":
        num /= 10
    return num


def _parse_g(value_name):
    if not value_name:
        return None
    m = re.search(r"([\d.,]+)\s*(kg|g)\b", value_name, re.IGNORECASE)
    if not m:
        return None
    num = float(m.group(1).replace(",", "."))
    return num * 1000 if m.group(2).lower() == "kg" else num


def fetch_ml_catalog_dimensions(session, title, description=None):
    """Segunda fuente de medidas reales (despues de las de Rodo): busca una
    coincidencia de catalogo con la MISMA validacion estricta que
    find_confident_catalog_matches (marca + codigo de modelo especifico, no
    solo palabras sueltas) y lee HEIGHT/WIDTH/DEPTH/WEIGHT del producto.

    El catalogo de ML tambien puede tener datos mal cargados -- se
    encontro un caso real de un TV con '7.17 m' de ancho y '1.237 cm' de
    alto, claramente un error de carga de ML. Por eso se descarta
    cualquier valor fuera de un rango fisicamente razonable
    (DIMENSION_SANITY_RANGE_CM / WEIGHT_SANITY_RANGE_G) en vez de
    confiarlo ciegamente."""
    row_stub = {"title": title, "description": description or ""}
    try:
        candidates = find_confident_catalog_matches(session, row_stub)
    except Exception:
        return None
    if not candidates:
        return None

    resp = session.get(f"{API_BASE}/products/{candidates[0]}")
    if not resp.ok:
        return None
    attrs = {a.get("id"): a.get("value_name") for a in resp.json().get("attributes", [])}

    h = _parse_cm(attrs.get("HEIGHT") or attrs.get("PACKAGE_HEIGHT"))
    w = _parse_cm(attrs.get("WIDTH") or attrs.get("PACKAGE_WIDTH"))
    d = _parse_cm(attrs.get("DEPTH") or attrs.get("LENGTH") or attrs.get("PACKAGE_LENGTH"))
    wt = _parse_g(attrs.get("WEIGHT") or attrs.get("PACKAGE_WEIGHT"))

    lo, hi = DIMENSION_SANITY_RANGE_CM
    h = h if h and lo <= h <= hi else None
    w = w if w and lo <= w <= hi else None
    d = d if d and lo <= d <= hi else None
    wlo, whi = WEIGHT_SANITY_RANGE_G
    wt = wt if wt and wlo <= wt <= whi else None

    if h and w and d:
        return (round(h + PACKAGING_BUFFER_CM), round(w + PACKAGING_BUFFER_CM),
                round(d + PACKAGING_BUFFER_CM), round(wt) if wt else None)
    return None


def estimate_package(title, category_name, description=None, session=None):
    """Prioridad: 1) medidas reales de Rodo (mas el buffer de embalaje),
    2) medidas reales del catalogo de ML (si session esta disponible),
    3) palabra clave del titulo, 4) promedio por categoria, 5) default
    generico. El peso sale de la estimacion por palabra clave/categoria
    salvo que el catalogo de ML tenga un peso real y valido."""
    text = (title or "").lower()
    weight = None
    for keywords, dims in PACKAGE_KEYWORDS:
        if any(re.search(rf"\b{re.escape(k.strip())}\b", text) for k in keywords):
            weight = dims[3]
            break
    if weight is None:
        weight = CATEGORY_ESTIMATES.get((category_name or "").strip().lower(), PACKAGE_DEFAULT)[3]

    real_dims = parse_rodo_dimensions(description)
    if real_dims:
        height, width, length = real_dims
        return (height, width, length, weight)

    if session:
        catalog_dims = fetch_ml_catalog_dimensions(session, title, description)
        if catalog_dims:
            height, width, length, catalog_weight = catalog_dims
            return (height, width, length, catalog_weight or weight)

    for keywords, dims in PACKAGE_KEYWORDS:
        if any(re.search(rf"\b{re.escape(k.strip())}\b", text) for k in keywords):
            return dims
    return CATEGORY_ESTIMATES.get((category_name or "").strip().lower(), PACKAGE_DEFAULT)


def guess_brand(description):
    m = BRAND_ATTR_LINE.search(description or "")
    return m.group(1).strip() if m else None


def parse_description_attrs(description):
    """'Label: value' por linea -> {label en minuscula: value}."""
    attrs = {}
    for line in (description or "").splitlines():
        if ":" in line:
            label, _, value = line.partition(":")
            label, value = label.strip(), value.strip()
            if label and value:
                attrs[label.lower()] = value
    return attrs


def get_category_attributes(session, category_id):
    if category_id not in _category_attrs_cache:
        resp = session.get(f"{API_BASE}/categories/{category_id}/attributes")
        _category_attrs_cache[category_id] = resp.json() if resp.ok else []
    return _category_attrs_cache[category_id]


def find_attribute_id_by_name(session, category_id, attr_name):
    for attr in get_category_attributes(session, category_id):
        if attr.get("name", "").strip().lower() == attr_name.strip().lower():
            return attr["id"]
    return None


def find_attribute_name_by_id(session, category_id, attr_id):
    for attr in get_category_attributes(session, category_id):
        if attr.get("id") == attr_id:
            return attr.get("name")
    return None


def find_attribute_def_by_id(session, category_id, attr_id):
    for attr in get_category_attributes(session, category_id):
        if attr.get("id") == attr_id:
            return attr
    return None


NEUTRAL_VALUE_HINTS = ("sin genero", "unisex", "sin genero infantil")


def coerce_to_valid_value(attr_def, desired_value):
    """Si el atributo es de lista fija (value_type='list'), el texto libre
    no sirve: hay que mandar una de las opciones reales. Busca la opcion mas
    parecida al valor deseado; si no encuentra ninguna razonable, devuelve
    None (mejor no mandar nada a que rebote con un valor invalido)."""
    if not attr_def or attr_def.get("value_type") != "list":
        return desired_value

    options = attr_def.get("values") or []
    if not options:
        return None

    desired_norm = normalize_label(desired_value)
    for opt in options:
        if normalize_label(opt["name"]) == desired_norm:
            return opt["name"]
    for opt in options:
        if desired_norm in normalize_label(opt["name"]):
            return opt["name"]
    for hint in NEUTRAL_VALUE_HINTS:
        for opt in options:
            if hint in normalize_label(opt["name"]):
                return opt["name"]
    return None


def strip_accents(text):
    return "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")


STOPWORDS = {"de", "del", "la", "el", "los", "las", "un", "una", "en"}


def normalize_label(text):
    words = [w for w in strip_accents(text.lower()).split() if w not in STOPWORDS]
    return " ".join(words)


_GENERIC_DEFAULTS_NORM = {normalize_label(k): v for k, v in GENERIC_ATTRIBUTE_DEFAULTS.items()}


def resolve_attribute_value(attr_name, scraped_attrs, title):
    """Busca un valor razonable: primero en lo scrapeado, despues sinonimos
    conocidos, despues genericos, despues palabras clave del titulo."""
    name_norm = normalize_label(attr_name)

    for label, value in scraped_attrs.items():
        label_norm = normalize_label(label)
        if name_norm == label_norm or name_norm in label_norm or label_norm in name_norm:
            return value

    for alias in ATTRIBUTE_ALIASES.get(name_norm, []):
        for label, value in scraped_attrs.items():
            if alias in normalize_label(label):
                return value

    if name_norm in ("modelo", "titulo videojuego", "nombre producto"):
        return title[:60]

    if name_norm == "plataforma":
        title_low = strip_accents((title or "").lower())
        for keyword, value in PLATFORM_KEYWORDS:
            if keyword in title_low:
                return value

    from_title = resolve_from_title(name_norm, title)
    if from_title:
        return from_title

    if name_norm in _GENERIC_DEFAULTS_NORM:
        return _GENERIC_DEFAULTS_NORM[name_norm]
    return None


TITLE_NUMBER_PATTERNS = {
    "capacidad volumen": (re.compile(r"(\d+)\s*lts?\b", re.IGNORECASE), "{} L"),
    "capacidad neta": (re.compile(r"(\d+)\s*lts?\b", re.IGNORECASE), "{} L"),
    "capacidad botellas": (re.compile(r"(\d+)\s*botellas?\b", re.IGNORECASE), "{}"),
    "capacidad carga": (re.compile(r"(\d+)\s*[Kk]g\b"), "{} kg"),
    "velocidad centrifugado": (re.compile(r"(\d+)\s*rpm\b", re.IGNORECASE), "{} rpm"),
}


def resolve_from_title(name_norm, title):
    """Ultimo recurso antes de los genericos fijos: algunos atributos
    numericos vienen literalmente escritos en el titulo de Rodo (ej.
    '153Lts', '32 Botellas', '9Kg', '1400rpm') o se pueden inferir de una
    palabra presente/ausente (ej. 'C/Freezer', 'C/Superior')."""
    title_low = strip_accents((title or "").lower())

    if name_norm == "con freezer":
        if "sin freezer" in title_low:
            return "No"
        if "freezer" in title_low or "frizer" in title_low or "side by side" in title_low:
            return "Sí"
        if "bajo mesada" in title_low or "compacta" in title_low:
            return "No"  # convencion: bajo mesada/compacta suelen ser sin freezer separado
        return None

    if name_norm == "tipo lavarropas":
        if "c/superior" in title_low or "carga superior" in title_low:
            return "Carga Superior"
        if "c/frontal" in title_low or "carga frontal" in title_low:
            return "Carga Frontal"
        return None

    if name_norm == "tipos horno":
        if "electrico" in title_low:
            return "Eléctrico"
        if " a gas" in title_low or "gas natural" in title_low:
            return "A Gas"
        return None

    if name_norm == "lugares montaje":
        if "soporte" in title_low:
            return "Pared"  # los "Soporte TV" de Rodo son soportes de pared
        return "Mesa"  # audio/parlantes standalone: se apoyan en mesa/piso

    if name_norm == "tipo montaje":
        if "empotrar" in title_low or "empotrable" in title_low or "empotrado" in title_low:
            return "Empotrable"
        if "independiente" in title_low:
            return "Independiente"
        return None

    pattern_unit = TITLE_NUMBER_PATTERNS.get(name_norm)
    if pattern_unit:
        pattern, template = pattern_unit
        m = pattern.search(title or "")
        if m:
            return template.format(m.group(1))
    return None


def suggest_category(session, site, title):
    resp = session.get(f"{API_BASE}/sites/{site}/domain_discovery/search",
                        params={"q": title[:60], "limit": 1})
    if resp.ok and resp.json():
        return resp.json()[0]["category_id"]
    return None


MODEL_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9\-/]*\d[A-Za-z0-9\-/]*")


def extract_model_tokens(title):
    """Codigos de modelo (contienen letras Y numeros, ej 'MSNIC-18H-GN81F',
    'L3350'). Son mucho mas especificos que palabras sueltas para confirmar
    que una ficha de catalogo es realmente el mismo producto."""
    tokens = set()
    for m in MODEL_TOKEN_RE.finditer(title or ""):
        tok = strip_accents(m.group(0).lower())
        if len(tok) >= 3:
            tokens.add(tok)
    return tokens


def fetch_catalog_product_name(session, product_id):
    resp = session.get(f"{API_BASE}/products/{product_id}")
    if not resp.ok:
        return None
    return resp.json().get("name")


def find_confident_catalog_matches(session, row):
    """Busca en el catalogo de ML una variante que matchee con CONFIANZA
    ALTA: exige que el codigo de modelo mas especifico de nuestro titulo
    (ej 'MSNIC-18H-GN81F') aparezca literal en el nombre de la variante
    puntual (no el generico "padre"), y que la marca tambien coincida.
    Sin esto, en pruebas reales enganchaba a marcas y bundles equivocados
    (ver MANUAL.md seccion 8) -- esta version es deliberadamente estricta:
    prefiere no encontrar match antes que encontrar uno dudoso."""
    title = row.get("title") or ""
    brand = strip_accents((guess_brand(row.get("description")) or "").lower())
    model_tokens = extract_model_tokens(title)

    if not model_tokens:
        return []  # sin codigo de modelo identificable, no arriesgamos

    primary_token = max(model_tokens, key=len)
    if len(primary_token) < 4:
        return []  # muy corto/generico (ej. solo "630") para confiar solo

    resp = session.get(f"{API_BASE}/products/search", params={"site_id": "MLA", "q": title[:60]})
    if not resp.ok:
        return []

    candidates = []
    for result in resp.json().get("results", [])[:5]:
        children = result.get("children_ids") or []
        candidate_id = children[0] if children else result["id"]
        candidate_name = (fetch_catalog_product_name(session, candidate_id)
                           if children else result.get("name", ""))
        if not candidate_name:
            continue
        name_norm = strip_accents(candidate_name.lower())

        if primary_token not in name_norm:
            continue
        if brand and brand not in name_norm:
            continue
        name_tokens = set(name_norm.split())
        our_tokens = set(strip_accents(title.lower()).split())
        extra = {t for t in name_tokens if t not in our_tokens and len(t) > 2}
        if extra & BUNDLE_RISK_WORDS:
            continue

        candidates.append(candidate_id)
    return candidates


IMAGE_TIMEOUT = 10
MIN_COVER_SIDE = 400
GOOD_ASPECT_RANGE = (0.7, 1.4)


def fetch_image(url):
    """Descarga la imagen completa (para leer dimensiones y analizar esquinas)."""
    try:
        resp = requests.get(url, timeout=IMAGE_TIMEOUT)
        resp.raise_for_status()
        return Image.open(io.BytesIO(resp.content))
    except Exception:
        return None


def get_image_size(url):
    img = fetch_image(url)
    return img.size if img else None


BADGE_CORNER_FRAC = 0.35


def corner_badge_score(img):
    """Puntaje HEURISTICO (0 = mas limpio) de si alguna esquina de la
    imagen tiene un sticker/banner promocional superpuesto (ej. "3 CUOTAS
    SIN INTERES", circulos de "ENVIO SAME DAY" que trae Rodo en las fotos
    de algunos productos). Se probo esto como filtro duro y NO es confiable
    -- un producto genuinamente colorido cerca de un borde (ej. una pava
    celeste) puntua igual o mas alto que un banner real, asi que NO usar
    para descartar fotos. Sirve solo como DESEMPATE de baja confianza entre
    fotos que ya son candidatas validas por tamano/aspecto: ante la duda,
    preferir la de puntaje mas bajo, nunca excluir una foto solo por esto."""
    best = 0.0
    w, h = img.size
    cw, ch = int(w * BADGE_CORNER_FRAC), int(h * BADGE_CORNER_FRAC)
    boxes = {
        "tl": (0, 0, cw, ch), "tr": (w - cw, 0, w, ch),
        "bl": (0, h - ch, cw, h), "br": (w - cw, h - ch, w, h),
    }
    rgb = img.convert("RGB")
    for box in boxes.values():
        crop = rgb.crop(box)
        pixels = list(crop.getdata())
        n = len(pixels) or 1
        sat_frac = sum(1 for p in pixels if (max(p) - min(p)) > 60 and max(p) > 100) / n
        if sat_frac < 0.03:
            continue
        edges = crop.convert("L").filter(ImageFilter.FIND_EDGES)
        edge_frac = sum(1 for p in edges.getdata() if p > 40) / n
        best = max(best, sat_frac * edge_frac)
    return best


def pick_cover_first(urls, max_check=8):
    """Reordena la galeria de fotos para que la portada sea razonablemente
    cuadrada. ML rechaza como portada fotos muy panoramicas/angostas (ej.
    cortes tecnicos horizontales de 500x120) o muy chicas -- eso genero
    publicaciones 'under_review: forbidden' en produccion. Solo revisa las
    primeras `max_check` fotos para no demorar demasiado por producto.

    max_check subido de 4 a 8: con 4, productos con forma inherentemente
    ancha/chata (purificadores de cocina, toalleros) a veces tienen sus
    primeras 4 fotos todas panoramicas y ninguna pasa el filtro -- mirar
    mas fotos de la galeria aumenta la chance de encontrar un angulo mejor
    (ej. de 3/4) mas atras en el orden original. Si NINGUNA de las 8 pasa,
    sigue sin haber garantia de portada valida -- estos productos pueden
    necesitar revision manual igual.

    Entre las candidatas "buenas" por tamano/aspecto, desempata con
    corner_badge_score (ver docstring) para preferir, cuando hay opcion,
    la que parece no tener un sticker de cuotas/envio superpuesto."""
    if not urls:
        return urls

    scored = []
    for i, u in enumerate(urls):
        if i >= max_check:
            scored.append((1, 0.0, i, u))  # no revisada, prioridad media, mantiene orden
            continue
        img = fetch_image(u)
        if not img:
            scored.append((1, 0.0, i, u))
            continue
        w, h = img.size
        if w < MIN_COVER_SIDE or h < MIN_COVER_SIDE:
            scored.append((2, 0.0, i, u))  # muy chica, mala portada
        elif GOOD_ASPECT_RANGE[0] <= (w / h) <= GOOD_ASPECT_RANGE[1]:
            scored.append((0, corner_badge_score(img), i, u))  # buena portada
        else:
            scored.append((2, 0.0, i, u))  # panoramica/angosta, evitar como portada

    scored.sort(key=lambda x: (x[0], x[1], x[2]))
    return [u for _, _, _, u in scored]


_ML_FEE_CACHE = {}


def get_ml_fee_info(session, site, category_id, listing_type, price_estimate):
    """Consulta la comisión real de ML para esa categoria+tipo+precio.
    Devuelve (porcentaje_fee, fee_fijo). Se cachea por (categoria, tipo,
    banda de precio) porque la comisión puede variar por rango de precio."""
    band = round(price_estimate, -3) if price_estimate else 0  # cachear por banda de $1000
    cache_key = (category_id, listing_type, band)
    if cache_key in _ML_FEE_CACHE:
        return _ML_FEE_CACHE[cache_key]

    resp = session.get(f"{API_BASE}/sites/{site}/listing_prices",
                        params={"price": price_estimate, "listing_type_id": listing_type,
                                "category_id": category_id})
    if not resp.ok:
        result = (0.0, 0.0)
    else:
        details = resp.json().get("sale_fee_details", {})
        result = (details.get("percentage_fee", 0.0) or 0.0, details.get("fixed_fee", 0.0) or 0.0)
    _ML_FEE_CACHE[cache_key] = result
    return result


def compute_sale_price(session, site, price_original, category_id, listing_type, markup, account_for_ml_fee=False):
    """Precio de venta = precio_original * markup (ej. markup=1.8 -> +80%).

    Por default es un calculo simple, SIN mirar la comision de venta de ML
    (asi lo pidio el vendedor: con un margen bruto del 80% ya sobra de
    sobra para cubrir la comision de ML, que ronda 10-16%). Si se pasa
    account_for_ml_fee=True (--account-for-ml-fee), en cambio calcula el
    precio para que `markup` quede como margen NETO ya descontada la
    comision real de ML por categoria (mas preciso, pero mas dificil de
    verificar a ojo) -- ver MANUAL.md seccion 8, ahi se detecto vendiendo
    a perdida real usando el calculo simple con un margen muy chico (1.8%),
    que no alcanzaba a cubrir la comision de ML. Con un markup grande como
    1.8 (80%) ese riesgo ya no aplica."""
    if not price_original:
        return None
    if not account_for_ml_fee:
        return round(price_original * markup, 2)

    price = price_original * markup
    for _ in range(3):
        pct, fixed = get_ml_fee_info(session, site, category_id, listing_type, price)
        new_price = (price_original * markup + fixed) / (1 - pct / 100) if pct < 100 else price
        if abs(new_price - price) < 1:
            price = new_price
            break
        price = new_price
    return round(price, 2)


def build_item_payload(session, row, category_id, listing_type, condition, markup_chico, markup_grande,
                        status="paused", site="MLA", account_for_ml_fee=False):
    title = (row.get("title") or "")[:60]
    image_urls = pick_cover_first([u for u in (row.get("images") or "").split("|") if u])
    pictures = [{"source": u} for u in image_urls]
    height, width, length, weight = estimate_package(row.get("title"), row.get("category"), row.get("description"), session=session)

    attributes = [
        {"id": "GTIN", "value_name": "N/A"},
        {"id": "SELLER_PACKAGE_HEIGHT", "value_name": f"{height} cm"},
        {"id": "SELLER_PACKAGE_WIDTH", "value_name": f"{width} cm"},
        {"id": "SELLER_PACKAGE_LENGTH", "value_name": f"{length} cm"},
        {"id": "SELLER_PACKAGE_WEIGHT", "value_name": f"{weight} g"},
    ]
    brand = guess_brand(row.get("description"))
    if brand:
        attributes.append({"id": "BRAND", "value_name": brand})
    if row.get("sku"):
        # SELLER_SKU es el atributo que llena el panel "Codigo de
        # identificacion (SKU)" en el editor de ML -- es DISTINTO de
        # seller_custom_field (mas abajo), que es un campo legacy que no
        # se muestra ahi. Mandamos los dos por las dudas.
        attributes.append({"id": "SELLER_SKU", "value_name": str(row["sku"])})

    markup = pick_markup(row.get("title"), markup_chico, markup_grande)
    price = compute_sale_price(session, site, row.get("price_original"), category_id, listing_type, markup,
                                account_for_ml_fee=account_for_ml_fee)

    payload = {
        "family_name": title,
        "category_id": category_id,
        "price": price,
        "currency_id": "ARS",
        "available_quantity": DEFAULT_STOCK,
        "buying_mode": "buy_it_now",
        "condition": condition,
        "listing_type_id": listing_type,
        "status": status,
        "pictures": pictures,
        "attributes": attributes,
    }
    if row.get("sku"):
        payload["seller_custom_field"] = row["sku"]
    return payload


def build_catalog_payload(session, row, category_id, catalog_product_id, listing_type, condition, markup_chico,
                           markup_grande, status="paused", site="MLA", account_for_ml_fee=False):
    height, width, length, weight = estimate_package(row.get("title"), row.get("category"), row.get("description"), session=session)
    markup = pick_markup(row.get("title"), markup_chico, markup_grande)
    price = compute_sale_price(session, site, row.get("price_original"), category_id, listing_type, markup,
                                account_for_ml_fee=account_for_ml_fee)
    payload = {
        "catalog_product_id": catalog_product_id,
        "catalog_listing": True,
        "category_id": category_id,
        "price": price,
        "currency_id": "ARS",
        "available_quantity": DEFAULT_STOCK,
        "buying_mode": "buy_it_now",
        "condition": condition,
        "listing_type_id": listing_type,
        "status": status,
        "attributes": [
            {"id": "SELLER_PACKAGE_HEIGHT", "value_name": f"{height} cm"},
            {"id": "SELLER_PACKAGE_WIDTH", "value_name": f"{width} cm"},
            {"id": "SELLER_PACKAGE_LENGTH", "value_name": f"{length} cm"},
            {"id": "SELLER_PACKAGE_WEIGHT", "value_name": f"{weight} g"},
        ] + ([{"id": "SELLER_SKU", "value_name": str(row["sku"])}] if row.get("sku") else []),
    }
    if row.get("sku"):
        payload["seller_custom_field"] = row["sku"]
    return payload


def create_item_with_retries(session, payload, category_id, row):
    """Publica, y si ML pide un atributo puntual de catalogo, lo completa y reintenta.

    ML a veces rechaza el GTIN="N/A" (cuando el producto le resulta parecido a
    uno de su catalogo) y a veces no le importa, incluso para el mismo
    producto en llamadas distintas -- no encontramos una regla 100%
    determinista. Como ultimo recurso, si choca especificamente con el
    formato de GTIN, se saca ese atributo del todo y se reintenta una vez
    (sin el, ML a veces deja de exigir el match de catalogo)."""
    scraped_attrs = parse_description_attrs(row.get("description"))
    seen_attr_ids = {a["id"] for a in payload["attributes"]}
    dropped_gtin = False

    resp = None
    for _ in range(MAX_ATTRIBUTE_RETRIES):
        resp = session.post(f"{API_BASE}/items", json=payload)
        if resp.ok:
            return resp

        try:
            causes = resp.json().get("cause", [])
        except ValueError:
            causes = []

        if not dropped_gtin and any(GTIN_INVALID_MSG.search(c.get("message", "")) for c in causes):
            payload["attributes"] = [a for a in payload["attributes"] if a["id"] != "GTIN"]
            seen_attr_ids.discard("GTIN")
            dropped_gtin = True
            continue

        progress = False

        invalid_ids = {m.group(1) for c in causes
                       for m in [INVALID_ATTR_ID_MSG.search(c.get("message", ""))] if m}
        for attr_id in invalid_ids:
            entry = next((a for a in payload["attributes"] if a["id"] == attr_id), None)
            if not entry:
                continue
            attr_def = find_attribute_def_by_id(session, category_id, attr_id)
            fixed = coerce_to_valid_value(attr_def, entry["value_name"])
            if fixed and fixed != entry["value_name"]:
                entry["value_name"] = fixed
                progress = True
            else:
                payload["attributes"] = [a for a in payload["attributes"] if a["id"] != attr_id]
                seen_attr_ids.discard(attr_id)

        missing_pairs = []  # (attr_id or None, attr_name)
        for c in causes:
            message = c.get("message", "")
            m = MISSING_ATTR_MSG.search(message)
            if m:
                missing_pairs.append((None, m.group(1)))
            m2 = MISSING_ATTR_IDS_MSG.search(message)
            if m2:
                for attr_id in [x.strip() for x in m2.group(1).split(",") if x.strip()]:
                    missing_pairs.append((attr_id, None))

        for attr_id, attr_name in missing_pairs:
            if not attr_id:
                attr_id = find_attribute_id_by_name(session, category_id, attr_name)
            if not attr_name:
                attr_name = find_attribute_name_by_id(session, category_id, attr_id)
            if not attr_id or not attr_name or attr_id in seen_attr_ids:
                continue
            value = resolve_attribute_value(attr_name, scraped_attrs, row.get("title") or "")
            if not value:
                continue
            attr_def = find_attribute_def_by_id(session, category_id, attr_id)
            value = coerce_to_valid_value(attr_def, value)
            if not value:
                continue
            payload["attributes"].append({"id": attr_id, "value_name": value[:60]})
            seen_attr_ids.add(attr_id)
            progress = True

        if not invalid_ids and not missing_pairs:
            return resp

        if not progress:
            return resp

    return resp


def force_status(session, item_id, desired_status, retries=6, delay=5):
    """Confirma el estado final con un PUT explicito DESPUES de crear el
    item (el mismo tipo de accion que haria un vendedor a mano en ML).

    El campo "status" que mandamos al CREAR no alcanza por si solo: en
    produccion se observaron publicaciones que ML paso de 'paused' a
    'active' por su cuenta un rato despues de creadas (aparentemente al
    terminar de procesar las fotos), sin que nosotros lo pidieramos. La
    confirmacion por PUT demostro sostenerse. NUNCA confiar solo en el
    status de creacion para items reales -- siempre confirmar con esto
    despues (main() ya lo hace automaticamente).

    OJO: pasar de 'paused' a 'active' inmediatamente despues de crear el
    item puede fallar las primeras veces porque ML todavia esta
    procesando la publicacion nueva (fotos, indexado) -- se detecto en
    produccion que reintentando unos segundos despues SI funciona con los
    mismos datos. Por eso son 6 reintentos con 5 segundos de espera (30s
    en total) en vez de 3 reintentos de 1 segundo."""
    last_resp = None
    for attempt in range(retries):
        resp = session.put(f"{API_BASE}/items/{item_id}", json={"status": desired_status})
        last_resp = resp
        if resp.ok and resp.json().get("status") == desired_status:
            return True
        if resp.ok and resp.json().get("status") == "under_review":
            return True  # bajo revision de ML tampoco es publico; no se puede forzar mas
        time.sleep(delay)
    if last_resp is not None:
        print(f"    (no se pudo confirmar '{desired_status}' para {item_id} tras {retries} intentos: "
              f"{last_resp.status_code} {last_resp.text[:200]})")
    return False


def is_gtin_wall(resp):
    if resp.ok:
        return False
    try:
        causes = resp.json().get("cause", [])
    except ValueError:
        return False
    return any(GTIN_INVALID_MSG.search(c.get("message", "")) for c in causes)


def create_item(session, row, category_id, listing_type, condition, markup_chico, markup_grande, site="MLA",
                 allow_catalog_attach=False, account_for_ml_fee=False):
    """Crea el item SIEMPRE en estado 'paused' primero (borrador). Esto
    evita el muro de GTIN real: MercadoLibre valida ese requisito de forma
    mucho mas estricta para items activos que para borradores recien
    creados. Una vez creado, main() confirma el estado final que se haya
    pedido (paused o active) con un PUT explicito por separado -- ver
    force_status() y la seccion 8 de MANUAL.md (ese PUT de confirmacion es
    el que demostro sostenerse; el status de la creacion por si solo no).

    Si por algun atributo puntual (Coleccion, Voltaje, etc.) ML igual no
    deja crear ni el borrador, create_item_with_retries ya intento
    completarlo con lo scrapeado o con un valor generico razonable; si
    ni asi alcanza, el producto queda como ERROR para revisar a mano.

    IMPORTANTE: hubo una version anterior que intentaba enganchar
    automaticamente a una variante de catalogo (--allow-catalog-attach) para
    sortear el GTIN. Se detecto que elegia variantes equivocadas (marca
    distinta, bundles con regalo que Rodo no incluye, o hasta el
    repuesto/control en vez del producto principal) en una fraccion no
    despreciable de casos, generando publicaciones reales incorrectas.
    Ahora exige coincidencia estricta de marca + codigo de modelo especifico
    (ver find_confident_catalog_matches), pero sigue siendo mas riesgoso
    que el flujo directo -- revisar MANUAL.md seccion 8 antes de usarlo."""
    payload = build_item_payload(session, row, category_id, listing_type, condition, markup_chico, markup_grande,
                                  site=site, account_for_ml_fee=account_for_ml_fee)
    resp = create_item_with_retries(session, payload, category_id, row)

    if allow_catalog_attach and is_gtin_wall(resp):
        catalog_ids = find_confident_catalog_matches(session, row)
        for catalog_id in catalog_ids:
            catalog_payload = build_catalog_payload(session, row, category_id, catalog_id, listing_type, condition,
                                                      markup_chico, markup_grande, site=site,
                                                      account_for_ml_fee=account_for_ml_fee)
            catalog_resp = session.post(f"{API_BASE}/items", json=catalog_payload)
            if catalog_resp.ok:
                return catalog_resp, "catalogo"
            resp = catalog_resp
        if catalog_ids:
            return resp, "catalogo_fallido"

    return resp, "directo"


def load_published_state():
    """SKU -> item_id de todo lo ya publicado, para no duplicar en reintentos.
    PUBLISHED_STATE_FILE es la UNICA fuente de verdad: se actualiza sola
    (save_published_state) cada vez que una publicacion nueva se crea bien.

    OJO: antes esta funcion tambien releia los excels resultado_*.xlsx para
    "completar" el estado -- eso resucitaba SKUs que habiamos sacado a mano
    del json (por ejemplo, despues de pausar publicaciones mal enganchadas a
    catalogo) porque el excel viejo los seguia marcando como OK. Si alguna
    vez haces limpieza manual del json, ya no hay riesgo de que vuelva solo."""
    state = {}
    if os.path.exists(PUBLISHED_STATE_FILE):
        with open(PUBLISHED_STATE_FILE, encoding="utf-8") as f:
            state = json.load(f)
    return state


def save_published_state(state):
    with open(PUBLISHED_STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)


def read_rows(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def export_results(results, path):
    wb = Workbook()
    ws = wb.active
    headers = ["url", "sku", "title", "price_ml", "status", "detail", "metodo"]
    ws.append(headers)
    for r in results:
        ws.append([r.get(h) for h in headers])
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="infile", default="productos_rodo.xlsx")
    parser.add_argument("--out", dest="outfile", default="resultado_publicacion.xlsx")
    parser.add_argument("--site", default="MLA")
    parser.add_argument("--listing-type", default="gold_special",
                         help="Ver tipos disponibles para tu cuenta en"
                              " /sites/{site}/listing_types")
    parser.add_argument("--condition", default="new")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--dry-run", action="store_true",
                         help="Arma el payload y sugiere categoria pero no publica nada")
    parser.add_argument("--allow-catalog-attach", action="store_true",
                         help="Intenta enganchar a una variante de catalogo de ML cuando"
                              " exige GTIN real, exigiendo coincidencia estricta de marca +"
                              " codigo de modelo especifico (no solo palabras sueltas)."
                              " Aun asi puede fallar en casos raros -- revisar MANUAL.md"
                              " seccion 8 antes de usarlo a gran escala.")
    parser.add_argument("--publish-status", choices=["paused", "active"], default="paused",
                         help="Estado final de la publicacion despues de confirmada (default:"
                              " paused, para revisar antes de que sea publica). 'active' la"
                              " deja visible y comprable de inmediato -- usar con criterio.")
    parser.add_argument("--markup-chico", type=float, default=1.8,
                         help="Multiplicador de precio para articulos chicos (todo Electrohogar"
                              " excepto la lista de LARGE_ITEM_KEYWORDS). 1.8 = +80%% (default).")
    parser.add_argument("--markup-grande", type=float, default=2.0,
                         help="Multiplicador de precio para electrodomesticos grandes (heladera,"
                              " lavarropas, lavavajillas, aire acondicionado, cocina, secarropas,"
                              " freezer). 2.0 = +100%% (default). Ambos son calculos simples"
                              " (precio_rodo * markup), SIN mirar la comision de venta de ML."
                              " Ver --account-for-ml-fee si se quiere un margen garantizado neto.")
    parser.add_argument("--account-for-ml-fee", action="store_true",
                         help="Calcula el precio para que --markup quede como margen NETO"
                              " ya descontada la comision real de ML por categoria (mas"
                              " preciso, pero mas dificil de verificar a ojo). Usar sobre todo"
                              " si se baja --markup a un numero chico -- ver MANUAL.md seccion"
                              " 8 (se detecto vendiendo a perdida real con un margen chico sin"
                              " esto activado).")
    args = parser.parse_args()

    rows = read_rows(args.infile)
    if args.limit:
        rows = rows[: args.limit]

    published = load_published_state() if not args.dry_run else {}

    token = get_valid_token()
    session = TimeoutSession()
    session.headers.update({"Authorization": f"Bearer {token}"})

    results = []
    for i, row in enumerate(rows, 1):
      try:
        title = row.get("title") or ""
        sku = str(row.get("sku") or "")
        if not title or not row.get("price_original"):
            results.append({**row, "status": "SKIPPED", "detail": "falta titulo o precio", "metodo": ""})
            print(f"[{i}/{len(rows)}] SKIP {title[:50]} (sin titulo/precio)")
            continue

        if sku in published:
            results.append({**row, "status": "SKIPPED", "detail": f"ya publicado -> {published[sku]}", "metodo": ""})
            print(f"[{i}/{len(rows)}] SKIP {title[:50]} (ya publicado -> {published[sku]})")
            continue

        category_id = suggest_category(session, args.site, title)
        if not category_id:
            results.append({**row, "status": "ERROR", "detail": "no se pudo sugerir categoria", "metodo": ""})
            print(f"[{i}/{len(rows)}] ERROR {title[:50]} (sin categoria)")
            continue

        if args.dry_run:
            dims = estimate_package(row.get("title"), row.get("category"), row.get("description"), session=session)
            markup = pick_markup(row.get("title"), args.markup_chico, args.markup_grande)
            price = compute_sale_price(session, args.site, row.get("price_original"), category_id,
                                        args.listing_type, markup,
                                        account_for_ml_fee=args.account_for_ml_fee)
            results.append({**row, "status": "DRY_RUN",
                             "detail": f"{category_id} | markup x{markup} | precio venta ${price} | paquete est. {dims}",
                             "metodo": ""})
            print(f"[{i}/{len(rows)}] DRY-RUN {title[:50]} -> categoria {category_id} | markup x{markup} | "
                  f"costo Rodo ${row.get('price_original')} -> precio venta ${price} | paquete est. {dims}")
            continue

        resp, metodo = create_item(session, row, category_id, args.listing_type, args.condition,
                                    args.markup_chico, args.markup_grande, site=args.site,
                                    account_for_ml_fee=args.account_for_ml_fee,
                                    allow_catalog_attach=args.allow_catalog_attach)
        if resp.ok:
            item_id = resp.json()["id"]
            if metodo == "directo":
                session.post(f"{API_BASE}/items/{item_id}/description",
                             json={"plain_text": row.get("description") or ""})
            confirmed = force_status(session, item_id, args.publish_status)
            note = "" if confirmed else f" [OJO: no se pudo confirmar estado '{args.publish_status}', revisar a mano YA]"
            results.append({**row, "status": "OK", "detail": item_id, "metodo": metodo})
            print(f"[{i}/{len(rows)}] OK ({metodo}, {args.publish_status}) {title[:50]} -> {item_id}{note}")
            published[sku] = item_id
            save_published_state(published)
        else:
            results.append({**row, "status": "ERROR", "detail": resp.text[:400], "metodo": metodo})
            print(f"[{i}/{len(rows)}] ERROR {title[:50]}: {resp.status_code} {resp.text[:300]}")

        time.sleep(0.5)
      except requests.exceptions.RequestException as e:
        # un corte/timeout de red no debe tumbar toda la tanda -- se marca
        # este producto como ERROR y se sigue con el siguiente.
        results.append({**row, "status": "ERROR", "detail": f"error de red: {e}", "metodo": ""})
        print(f"[{i}/{len(rows)}] ERROR DE RED {row.get('title', '')[:50]}: {e}")
        time.sleep(2)

    export_results(results, args.outfile)
    print(f"\nListo. Resultado en {args.outfile}")


if __name__ == "__main__":
    main()
