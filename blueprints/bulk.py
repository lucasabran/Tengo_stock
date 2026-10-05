"""Cargas masivas desde Excel/CSV: ventas, gastos e ingresos de stock."""
import io
from datetime import date

from flask import Blueprint, Response, jsonify, request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import auth
from db import get_db, now_iso
from helpers import (
    ACCOUNT_PAYMENT_METHOD,
    normalize_currency,
    parse_csv_file,
    parse_date,
    parse_number,
    parse_xlsx_file,
)

bp = Blueprint("bulk", __name__)

MAX_ROWS = 5000
MAX_ERRORS_SHOWN = 40

STOCK_REASONS = {
    "compra_proveedor": "compra_proveedor",
    "compra": "compra_proveedor",
    "proveedor": "compra_proveedor",
    "ajuste_inventario": "ajuste_inventario",
    "inventario": "ajuste_inventario",
    "ajuste_manual": "ajuste_manual",
    "ajuste": "ajuste_manual",
    "devolucion_proveedor": "devolucion_proveedor",
    "otro": "otro",
}

EXPENSE_CATEGORIES = ["Alquiler", "Servicios", "Sueldos", "Mercaderia", "Impuestos", "Otro"]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def read_rows():
    """Devuelve (rows, error_response). Cada row es un dict con claves canonicas."""
    file = request.files.get("file")
    if not file or not file.filename:
        return None, (jsonify({"error": "no se recibio ningun archivo"}), 400)
    name = file.filename.lower()
    try:
        if name.endswith(".csv"):
            rows = parse_csv_file(file)
        elif name.endswith((".xlsx", ".xlsm")):
            rows = parse_xlsx_file(file)
        else:
            return None, (jsonify({"error": "formato no soportado, usa .xlsx o .csv"}), 400)
    except Exception as exc:  # archivo corrupto, codificacion rara, etc.
        return None, (jsonify({"error": f"no se pudo leer el archivo: {exc}"}), 400)
    if not rows:
        return None, (jsonify({"error": "el archivo no tiene filas para importar (revisa que los encabezados coincidan con la plantilla)"}), 400)
    if len(rows) > MAX_ROWS:
        return None, (jsonify({"error": f"el archivo tiene mas de {MAX_ROWS} filas, dividilo en partes"}), 400)
    return rows, None


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def template_response(filename, headers, notes, sheet_title="Datos"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Instrucciones")
    guide.column_dimensions["A"].width = 110
    for line in notes:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def summarize(created, errors, extra=None):
    body = {
        "created": created,
        "error_count": len(errors),
        "errors": errors[:MAX_ERRORS_SHOWN],
        "more_errors": max(len(errors) - MAX_ERRORS_SHOWN, 0),
    }
    if extra:
        body.update(extra)
    return jsonify(body)


# ---------------------------------------------------------------------------
# Ventas
# ---------------------------------------------------------------------------
@bp.route("/api/sales/template.xlsx")
def sales_template():
    return template_response(
        "plantilla_ventas.xlsx",
        ["venta", "fecha", "canal", "cliente", "sku", "cantidad", "precio", "moneda", "medio_de_pago", "descuento", "nota"],
        [
            "Plantilla de carga masiva de VENTAS",
            "",
            "Una fila por producto vendido. Si una venta tiene varios productos, repeti el mismo numero en la columna 'venta' (ej: 1, 1, 1).",
            "Si dejas 'venta' vacio, cada fila es una venta distinta.",
            "",
            "venta: numero o codigo para agrupar las filas de una misma venta (opcional).",
            "fecha: dd/mm/aaaa (opcional, si esta vacia se usa hoy).",
            "canal: nombre del canal (Mostrador, Mercado Libre, Tienda Nube, Instagram...). Vacio = Mostrador. Si no existe, se crea.",
            "cliente: nombre del cliente (opcional). Si no existe, se crea. Obligatorio si el medio de pago es 'Cuenta corriente'.",
            "sku: codigo del producto. Tiene que existir en Stock.",
            "cantidad: unidades vendidas (obligatorio).",
            "precio: precio unitario (opcional, vacio = precio actual del producto).",
            "moneda: ARS o USD (vacio = la moneda del producto).",
            "medio_de_pago: Efectivo, Transferencia, Tarjeta, Cuenta corriente, etc. (opcional).",
            "descuento: monto de descuento de toda la venta, en la primera fila de la venta (opcional).",
            "nota: comentario (opcional).",
            "",
            "Al importar podes elegir si se descuenta el stock de los productos (para ventas nuevas) o no (para historicas que el stock ya refleja).",
        ],
    )


@bp.route("/api/sales/import", methods=["POST"])
def import_sales():
    rows, err = read_rows()
    if err:
        return err
    deduct = request.form.get("deduct_stock", "1") not in ("0", "false", "no")

    db = get_db()
    errors, created_channels, created_customers = [], [], []
    created = 0

    # agrupar filas por venta, conservando el orden del archivo
    groups, order_keys = {}, []
    for pos, row in enumerate(rows, start=2):
        key = clean(row.get("order")) or f"__fila_{pos}"
        if key not in groups:
            groups[key] = []
            order_keys.append(key)
        groups[key].append((pos, row))

    now = now_iso()
    for key in order_keys:
        lines_in = groups[key]
        first_pos, first = lines_in[0]
        label = f"venta {key}" if not key.startswith("__fila_") else f"fila {first_pos}"

        try:
            sale_date = parse_date(first.get("date"))
        except ValueError as exc:
            errors.append(f"{label}: {exc}")
            continue
        created_at = f"{sale_date}T12:00:00" if sale_date else now

        channel_name = clean(first.get("channel")) or "Mostrador"
        customer_name = clean(first.get("customer"))
        payment_method = clean(first.get("payment_method"))
        note = clean(first.get("note"))
        if payment_method.lower() == ACCOUNT_PAYMENT_METHOD.lower():
            payment_method = ACCOUNT_PAYMENT_METHOD
            if not customer_name:
                errors.append(f"{label}: para cuenta corriente hay que indicar el cliente")
                continue
        try:
            discount = parse_number(first.get("discount")) if clean(first.get("discount")) else 0.0
        except ValueError:
            errors.append(f"{label}: descuento invalido")
            continue

        # validar productos y stock ANTES de escribir nada de esta venta
        lines, problems, needed = [], [], {}
        currency = None
        for pos, row in lines_in:
            sku = clean(row.get("sku"))
            if not sku:
                problems.append(f"fila {pos}: falta sku")
                continue
            try:
                qty = int(parse_number(row.get("quantity")))
            except (TypeError, ValueError):
                problems.append(f"fila {pos}: cantidad invalida")
                continue
            if qty <= 0:
                problems.append(f"fila {pos}: la cantidad debe ser mayor a 0")
                continue
            product = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
            if not product:
                problems.append(f"fila {pos}: el producto {sku} no existe en Stock")
                continue
            needed[sku] = needed.get(sku, 0) + qty
            if deduct and product["quantity"] < needed[sku]:
                problems.append(f"fila {pos}: {sku} tiene {product['quantity']} en stock y se piden {needed[sku]}")
                continue
            try:
                unit_price = parse_number(row.get("price")) if clean(row.get("price")) else product["price"]
            except ValueError:
                problems.append(f"fila {pos}: precio invalido")
                continue
            if currency is None:
                currency = normalize_currency(row.get("currency")) if clean(row.get("currency")) else (product["currency"] or "ARS")
            lines.append(
                {"sku": sku, "name": product["name"], "price": unit_price, "qty": qty, "total": round(unit_price * qty, 2)}
            )
        if problems:
            errors.append(f"{label}: " + "; ".join(problems))
            continue

        # canal y cliente (se crean si no existen)
        channel = db.execute("SELECT id FROM channels WHERE lower(name) = lower(?)", (channel_name,)).fetchone()
        if channel:
            channel_id = channel["id"]
        else:
            channel_id = db.execute(
                "INSERT INTO channels (name, type, external_account_id, active, created_at) VALUES (?, 'other', '', 1, ?)",
                (channel_name, now),
            ).lastrowid
            created_channels.append(channel_name)
        customer_id = None
        if customer_name:
            customer = db.execute("SELECT id FROM customers WHERE lower(name) = lower(?)", (customer_name,)).fetchone()
            if customer:
                customer_id = customer["id"]
            else:
                customer_id = db.execute(
                    "INSERT INTO customers (name, phone, email, doc_number, note, created_at, updated_at) "
                    "VALUES (?, '', '', '', '', ?, ?)",
                    (customer_name, now, now),
                ).lastrowid
                created_customers.append(customer_name)

        subtotal = round(sum(l["total"] for l in lines), 2)
        sale_id = db.execute(
            "INSERT INTO sales (channel_id, customer_id, status, subtotal, discount, total, payment_method, currency, note, "
            "created_at, created_by, stock_deducted) VALUES (?, ?, 'completed', ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (channel_id, customer_id, subtotal, discount, round(subtotal - discount, 2), payment_method, currency or "ARS",
             note, created_at, auth.actor(), 1 if deduct else 0),
        ).lastrowid
        for l in lines:
            db.execute(
                "INSERT INTO sale_items (sale_id, sku, product_name, unit_price, quantity, line_total) VALUES (?, ?, ?, ?, ?, ?)",
                (sale_id, l["sku"], l["name"], l["price"], l["qty"], l["total"]),
            )
            if deduct:
                db.execute("UPDATE products SET quantity = quantity - ?, updated_at = ? WHERE sku = ?", (l["qty"], now, l["sku"]))
                db.execute(
                    "INSERT INTO stock_movements (sku, change_qty, reason, reference_type, reference_id, created_at, created_by) "
                    "VALUES (?, ?, 'venta', 'sale', ?, ?, ?)",
                    (l["sku"], -l["qty"], sale_id, created_at, auth.actor()),
                )
        created += 1

    db.commit()
    auth.audit("ventas_importadas", "sale", "", f"{created} ventas, {len(errors)} con errores, stock {'descontado' if deduct else 'sin descontar'}")
    return summarize(created, errors, {"created_channels": created_channels, "created_customers": created_customers})


# ---------------------------------------------------------------------------
# Gastos
# ---------------------------------------------------------------------------
@bp.route("/api/expenses/template.xlsx")
def expenses_template():
    return template_response(
        "plantilla_gastos.xlsx",
        ["fecha", "categoria", "monto", "proveedor", "nota"],
        [
            "Plantilla de carga masiva de GASTOS",
            "",
            "Una fila por gasto.",
            "fecha: dd/mm/aaaa (opcional, vacia = hoy).",
            "categoria: " + ", ".join(EXPENSE_CATEGORIES) + ". Vacia = Otro.",
            "monto: importe del gasto en pesos (obligatorio).",
            "proveedor y nota: opcionales.",
        ],
    )


@bp.route("/api/expenses/import", methods=["POST"])
def import_expenses():
    rows, err = read_rows()
    if err:
        return err
    db = get_db()
    errors, created = [], 0
    now = now_iso()
    known = {c.lower(): c for c in EXPENSE_CATEGORIES}

    for pos, row in enumerate(rows, start=2):
        try:
            amount = parse_number(row.get("amount"))
            expense_date = parse_date(row.get("date")) or date.today().isoformat()
        except ValueError as exc:
            errors.append(f"fila {pos}: {exc}")
            continue
        if amount <= 0:
            errors.append(f"fila {pos}: el monto tiene que ser mayor a 0")
            continue
        raw_cat = clean(row.get("category")) or "Otro"
        category = known.get(raw_cat.lower(), raw_cat)
        db.execute(
            "INSERT INTO expenses (category, amount, expense_date, vendor, note, created_at, updated_at, created_by) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (category, amount, expense_date, clean(row.get("vendor")), clean(row.get("note")), now, now, auth.actor()),
        )
        created += 1

    db.commit()
    auth.audit("gastos_importados", "expense", "", f"{created} gastos, {len(errors)} con errores")
    return summarize(created, errors)


# ---------------------------------------------------------------------------
# Ingresos de stock
# ---------------------------------------------------------------------------
@bp.route("/api/stock-entries/template.xlsx")
def stock_entries_template():
    return template_response(
        "plantilla_ingresos_stock.xlsx",
        ["sku", "cantidad", "motivo", "nota"],
        [
            "Plantilla de carga masiva de INGRESOS DE STOCK",
            "",
            "Suma (o resta, con numero negativo) unidades a productos que YA existen en Stock.",
            "Para crear productos nuevos usa 'Importar stock' (con nombre, precio, categoria).",
            "",
            "sku: codigo del producto (obligatorio, tiene que existir).",
            "cantidad: unidades que ENTRAN. Con signo menos (-3) se descuentan.",
            "motivo: compra_proveedor, ajuste_inventario, ajuste_manual, devolucion_proveedor u otro. Vacio = compra_proveedor.",
            "nota: ej. 'Factura 123 - Proveedor X' (opcional).",
        ],
    )


@bp.route("/api/stock-entries/import", methods=["POST"])
def import_stock_entries():
    rows, err = read_rows()
    if err:
        return err
    db = get_db()
    errors, created = [], 0
    now = now_iso()

    for pos, row in enumerate(rows, start=2):
        sku = clean(row.get("sku"))
        if not sku:
            errors.append(f"fila {pos}: falta sku")
            continue
        try:
            qty = int(parse_number(row.get("quantity")))
        except (TypeError, ValueError):
            errors.append(f"fila {pos}: cantidad invalida")
            continue
        if qty == 0:
            errors.append(f"fila {pos}: la cantidad no puede ser 0")
            continue
        reason_raw = clean(row.get("reason")).lower().replace(" ", "_") or "compra_proveedor"
        reason = STOCK_REASONS.get(reason_raw)
        if not reason:
            errors.append(f"fila {pos}: motivo invalido ({reason_raw})")
            continue
        product = db.execute("SELECT quantity FROM products WHERE sku = ?", (sku,)).fetchone()
        if not product:
            errors.append(f"fila {pos}: el producto {sku} no existe en Stock")
            continue
        if product["quantity"] + qty < 0:
            errors.append(f"fila {pos}: {sku} tiene {product['quantity']} y no se pueden restar {-qty}")
            continue
        db.execute("UPDATE products SET quantity = quantity + ?, updated_at = ? WHERE sku = ?", (qty, now, sku))
        db.execute(
            "INSERT INTO stock_movements (sku, change_qty, reason, note, reference_type, reference_id, created_at, created_by) "
            "VALUES (?, ?, ?, ?, '', NULL, ?, ?)",
            (sku, qty, reason, clean(row.get("note")), now, auth.actor()),
        )
        created += 1

    db.commit()
    auth.audit("ingresos_stock_importados", "product", "", f"{created} movimientos, {len(errors)} con errores")
    return summarize(created, errors)
