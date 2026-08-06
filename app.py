import csv
import io
import os
import secrets
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path

from flask import Flask, Response, g, jsonify, render_template, request
from openpyxl import load_workbook

DB_PATH = Path(__file__).parent / "stock.db"

STOCK_USER = os.environ.get("STOCK_USER", "admin")
STOCK_PASSWORD = os.environ.get("STOCK_PASSWORD", "admin1234")

app = Flask(__name__)


@app.before_request
def require_login():
    auth = request.authorization
    valid = (
        auth is not None
        and secrets.compare_digest(auth.username, STOCK_USER)
        and secrets.compare_digest(auth.password, STOCK_PASSWORD)
    )
    if not valid:
        return Response(
            "Acceso restringido. Usuario y clave requeridos.",
            401,
            {"WWW-Authenticate": 'Basic realm="Stock"'},
        )

HEADER_MAP = {
    "sku": "sku",
    "codigo": "sku",
    "code": "sku",
    "name": "name",
    "nombre": "name",
    "producto": "name",
    "price": "price",
    "precio": "price",
    "quantity": "quantity",
    "cantidad": "quantity",
    "stock": "quantity",
    "description": "description",
    "descripcion": "description",
    "detalle": "description",
}


def normalize_header(text):
    text = (text or "").strip().lower()
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return text


def map_headers(headers):
    mapping = {}
    for idx, h in enumerate(headers):
        key = HEADER_MAP.get(normalize_header(h))
        if key:
            mapping[idx] = key
    return mapping


def parse_csv_file(file_storage):
    raw = file_storage.read().decode("utf-8-sig")

    best_rows, best_mapping = [], {}
    for delimiter in (",", ";", "\t"):
        rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
        if not rows:
            continue
        mapping = map_headers(rows[0])
        if len(mapping) > len(best_mapping):
            best_rows, best_mapping = rows, mapping

    items = []
    for row in best_rows[1:]:
        if not any(cell.strip() for cell in row):
            continue
        item = {}
        for idx, key in best_mapping.items():
            if idx < len(row):
                item[key] = row[idx]
        items.append(item)
    return items


def parse_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return float(text)


def parse_xlsx_file(file_storage):
    wb = load_workbook(filename=io.BytesIO(file_storage.read()), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    mapping = map_headers([str(h) if h is not None else "" for h in header])
    items = []
    for row in rows_iter:
        if row is None or all(cell is None for cell in row):
            continue
        item = {}
        for idx, key in mapping.items():
            if idx < len(row):
                item[key] = row[idx]
        items.append(item)
    return items


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS products (
            sku TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            description TEXT DEFAULT '',
            price REAL NOT NULL DEFAULT 0,
            quantity INTEGER NOT NULL DEFAULT 0,
            updated_at TEXT NOT NULL
        )
        """
    )
    db.commit()
    db.close()


def row_to_dict(row):
    return {
        "sku": row["sku"],
        "name": row["name"],
        "description": row["description"],
        "price": row["price"],
        "quantity": row["quantity"],
        "updated_at": row["updated_at"],
    }


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/products", methods=["GET"])
def list_products():
    q = request.args.get("q", "").strip()
    db = get_db()
    if q:
        like = f"%{q}%"
        rows = db.execute(
            "SELECT * FROM products WHERE sku LIKE ? OR name LIKE ? ORDER BY name",
            (like, like),
        ).fetchall()
    else:
        rows = db.execute("SELECT * FROM products ORDER BY name").fetchall()
    return jsonify([row_to_dict(r) for r in rows])


@app.route("/api/products", methods=["POST"])
def create_product():
    data = request.get_json(force=True) or {}
    sku = str(data.get("sku", "")).strip()
    name = str(data.get("name", "")).strip()
    if not sku or not name:
        return jsonify({"error": "sku y name son obligatorios"}), 400

    try:
        price = float(data.get("price", 0))
        quantity = int(data.get("quantity", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "price y quantity deben ser numericos"}), 400

    description = str(data.get("description", "")).strip()

    db = get_db()
    exists = db.execute("SELECT 1 FROM products WHERE sku = ?", (sku,)).fetchone()
    if exists:
        return jsonify({"error": f"ya existe un producto con sku {sku}"}), 409

    db.execute(
        "INSERT INTO products (sku, name, description, price, quantity, updated_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (sku, name, description, price, quantity, datetime.now().isoformat(timespec="seconds")),
    )
    db.commit()
    row = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
    return jsonify(row_to_dict(row)), 201


@app.route("/api/products/<sku>", methods=["PUT"])
def update_product(sku):
    db = get_db()
    row = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
    if not row:
        return jsonify({"error": "producto no encontrado"}), 404

    data = request.get_json(force=True) or {}
    name = str(data.get("name", row["name"])).strip()
    description = str(data.get("description", row["description"]))
    try:
        price = float(data.get("price", row["price"]))
        quantity = int(data.get("quantity", row["quantity"]))
    except (TypeError, ValueError):
        return jsonify({"error": "price y quantity deben ser numericos"}), 400

    db.execute(
        "UPDATE products SET name=?, description=?, price=?, quantity=?, updated_at=? WHERE sku=?",
        (name, description, price, quantity, datetime.now().isoformat(timespec="seconds"), sku),
    )
    db.commit()
    row = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
    return jsonify(row_to_dict(row))


@app.route("/api/products/<sku>/add-stock", methods=["POST"])
def add_stock(sku):
    data = request.get_json(force=True) or {}
    try:
        amount = int(data.get("amount", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "amount debe ser numerico"}), 400

    db = get_db()
    row = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
    if not row:
        return jsonify({"error": "producto no encontrado"}), 404

    new_qty = row["quantity"] + amount
    db.execute(
        "UPDATE products SET quantity=?, updated_at=? WHERE sku=?",
        (new_qty, datetime.now().isoformat(timespec="seconds"), sku),
    )
    db.commit()
    row = db.execute("SELECT * FROM products WHERE sku = ?", (sku,)).fetchone()
    return jsonify(row_to_dict(row))


@app.route("/api/products/template.csv")
def download_template():
    content = "sku,name,price,quantity,description\n"
    return Response(
        content,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=plantilla_stock.csv"},
    )


@app.route("/api/products/import", methods=["POST"])
def import_products():
    if "file" not in request.files or not request.files["file"].filename:
        return jsonify({"error": "no se recibio ningun archivo"}), 400

    file = request.files["file"]
    filename = file.filename.lower()

    try:
        if filename.endswith(".csv"):
            rows = parse_csv_file(file)
        elif filename.endswith(".xlsx") or filename.endswith(".xlsm"):
            rows = parse_xlsx_file(file)
        else:
            return jsonify({"error": "formato no soportado, usa .csv o .xlsx"}), 400
    except Exception as exc:
        return jsonify({"error": f"no se pudo leer el archivo: {exc}"}), 400

    db = get_db()
    created = 0
    updated = 0
    errors = []
    now = datetime.now().isoformat(timespec="seconds")

    for i, row in enumerate(rows, start=2):
        sku_raw = row.get("sku", "")
        if isinstance(sku_raw, float) and sku_raw.is_integer():
            sku_raw = int(sku_raw)
        sku = str(sku_raw or "").strip()
        name = str(row.get("name", "") or "").strip()
        if not sku or not name:
            errors.append(f"fila {i}: falta sku o nombre")
            continue
        try:
            price = parse_number(row.get("price", 0))
            quantity = int(parse_number(row.get("quantity", 0)))
        except (TypeError, ValueError):
            errors.append(f"fila {i}: precio o cantidad invalido")
            continue
        description = str(row.get("description", "") or "").strip()

        exists = db.execute("SELECT 1 FROM products WHERE sku = ?", (sku,)).fetchone()
        if exists:
            db.execute(
                "UPDATE products SET name=?, description=?, price=?, quantity=?, updated_at=? WHERE sku=?",
                (name, description, price, quantity, now, sku),
            )
            updated += 1
        else:
            db.execute(
                "INSERT INTO products (sku, name, description, price, quantity, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (sku, name, description, price, quantity, now),
            )
            created += 1

    db.commit()
    return jsonify({"created": created, "updated": updated, "errors": errors})


@app.route("/api/products/<sku>", methods=["DELETE"])
def delete_product(sku):
    db = get_db()
    db.execute("DELETE FROM products WHERE sku = ?", (sku,))
    db.commit()
    return jsonify({"ok": True})


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
