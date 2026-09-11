import csv
import io
import unicodedata

from openpyxl import load_workbook

HEADER_MAP = {
    "sku": "sku",
    "codigo": "sku",
    "code": "sku",
    "name": "name",
    "nombre": "name",
    "producto": "name",
    "price": "price",
    "precio": "price",
    "quantity": "quantity",
    "cantidad": "quantity",
    "stock": "quantity",
    "description": "description",
    "descripcion": "description",
    "detalle": "description",
    "category": "category",
    "categoria": "category",
}


def normalize_header(text):
    text = (text or "").strip().lower()
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return text


def map_headers(headers):
    mapping = {}
    for idx, h in enumerate(headers):
        key = HEADER_MAP.get(normalize_header(h))
        if key:
            mapping[idx] = key
    return mapping


def parse_csv_file(file_storage):
    raw = file_storage.read().decode("utf-8-sig")

    best_rows, best_mapping = [], {}
    for delimiter in (",", ";", "\t"):
        rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
        if not rows:
            continue
        mapping = map_headers(rows[0])
        if len(mapping) > len(best_mapping):
            best_rows, best_mapping = rows, mapping

    items = []
    for row in best_rows[1:]:
        if not any(cell.strip() for cell in row):
            continue
        item = {}
        for idx, key in best_mapping.items():
            if idx < len(row):
                item[key] = row[idx]
        items.append(item)
    return items


def build_csv(headers, rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue()


def parse_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return float(text)


def parse_xlsx_file(file_storage):
    wb = load_workbook(filename=io.BytesIO(file_storage.read()), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    mapping = map_headers([str(h) if h is not None else "" for h in header])
    items = []
    for row in rows_iter:
        if row is None or all(cell is None for cell in row):
            continue
        item = {}
        for idx, key in mapping.items():
            if idx < len(row):
                item[key] = row[idx]
        items.append(item)
    return items
